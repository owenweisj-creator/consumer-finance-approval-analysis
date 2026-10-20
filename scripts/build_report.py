from __future__ import annotations

import subprocess
import sys
from datetime import datetime
from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from src.charts import channel_bar, contribution_bar, monthly_trend, risk_scatter
from src.db import query_df

NAVY = RGBColor(0x17, 0x3B, 0x57)
TEAL = RGBColor(0x0E, 0x7C, 0x86)
GRAY = RGBColor(0x4B, 0x55, 0x63)
LIGHT = RGBColor(0x6B, 0x72, 0x80)


def fmt(value, digits=0):
    return f"{value:,.{digits}f}"


def pct(value, digits=2):
    return f"{value * 100:.{digits}f}%"


def pp(value, digits=2):
    return f"{value * 100:+.{digits}f}pp"


def font(run, size=10, bold=False, color=GRAY):
    run.font.name = "Microsoft YaHei"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = color


def body(doc, text, size=10, bold=False, color=GRAY, space_after=5):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(space_after)
    paragraph.paragraph_format.line_spacing = 1.35
    run = paragraph.add_run(text)
    font(run, size=size, bold=bold, color=color)
    return paragraph


def heading(doc, text, level=1):
    paragraph = doc.add_heading(text, level=level)
    for run in paragraph.runs:
        font(run, size=17 if level == 1 else 13, bold=True, color=NAVY if level == 1 else TEAL)
    return paragraph


def shade(cell, color):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), color)
    tc_pr.append(shd)


def table(doc, df: pd.DataFrame, widths=None):
    if df.empty:
        body(doc, "暂无数据。")
        return
    t = doc.add_table(rows=1, cols=len(df.columns))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, col in enumerate(df.columns):
        cell = t.rows[0].cells[i]
        cell.text = str(col)
        shade(cell, "173B57")
        for p in cell.paragraphs:
            for r in p.runs:
                font(r, size=8.5, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF))
    for _, row in df.iterrows():
        cells = t.add_row().cells
        for i, col in enumerate(df.columns):
            cells[i].text = str(row[col])
            for p in cells[i].paragraphs:
                for r in p.runs:
                    font(r, size=8.5)
    if widths:
        for row in t.rows:
            for i, w in enumerate(widths):
                row.cells[i].width = Inches(w)


def main() -> None:
    figures = ROOT / "outputs" / "figures"
    reports = ROOT / "outputs" / "report"
    figures.mkdir(parents=True, exist_ok=True)
    reports.mkdir(parents=True, exist_ok=True)

    kpi = query_df("SELECT * FROM `v_申请审批月度` ORDER BY `月份` DESC LIMIT 2")
    if len(kpi) < 2:
        raise RuntimeError("KPI 数据不足，至少需要两个月")
    cur = kpi.iloc[0]
    prev = kpi.iloc[1]
    trend = query_df("SELECT * FROM `v_申请审批月度` WHERE `月份` >= '2025-10' ORDER BY `月份`")
    channel = query_df("SELECT * FROM `v_渠道月度` WHERE `月份` = %s ORDER BY `进件量` DESC", [cur["月份"]])
    decomp = query_df("SELECT * FROM `v_渠道归因` ORDER BY `总效应`")
    risk = query_df("SELECT * FROM `v_还款风险`")
    risk_merged = channel.merge(risk, on="渠道类型", how="left")

    mix_total = float(decomp["结构效应"].sum())
    rate_total = float(decomp["比率效应"].sum())
    total_effect = float(decomp["总效应"].sum())
    app_growth = (float(cur["进件量"]) - float(prev["进件量"])) / float(prev["进件量"])
    approval_pp = (float(cur["审批通过率"]) - float(prev["审批通过率"])) * 100
    amount_growth = (float(cur["审批金额"]) - float(prev["审批金额"])) / float(prev["审批金额"])
    ticket_growth = (float(cur["件均"]) - float(prev["件均"])) / float(prev["件均"])

    monthly_trend(trend, figures / "01_月度趋势.png")
    channel_bar(channel, figures / "02_渠道进件与通过率.png")
    contribution_bar(decomp, figures / "03_渠道贡献.png")
    risk_scatter(risk_merged.dropna(subset=["DPD30率"]), figures / "04_通过率与风险.png")

    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)

    title = doc.add_heading(f"消费金融审批通过率下降专项分析（{cur['月份']}）", level=0)
    for run in title.runs:
        font(run, size=24, bold=True, color=NAVY)
    body(doc, f"数据锚点：2026-09-30｜生成时间：{datetime.now().strftime('%Y-%m-%d %H:%M')}｜数据源：Home Credit 公开数据", size=9, color=LIGHT)

    heading(doc, "一、摘要")
    body(
        doc,
        f"{cur['月份']} 进件量 {fmt(cur['进件量'])} 笔（环比 {app_growth:+.1%}），"
        f"审批通过率 {pct(cur['审批通过率'])}（环比 {approval_pp:+.1f} 个百分点）。"
        f"渠道分解显示，结构效应 {mix_total*100:+.2f}pp、比率效应 {rate_total*100:+.2f}pp、"
        f"总效应 {total_effect*100:+.2f}pp。通过率下降主要由低通过率渠道占比上升、"
        f"高通过率渠道占比收缩造成，同时各渠道自身也略有收紧。",
        size=10.5,
    )
    body(
        doc,
        f"审批金额 {fmt(cur['审批金额'])} 万元（环比 {amount_growth:+.1%}），但审批件数环比下降，"
        f"金额增长主要来自件均 {fmt(cur['件均'], 2)} 万元（环比 {ticket_growth:+.1%}）和期限拉长，"
        f"不是件数扩张。",
    )

    heading(doc, "二、背景与问题")
    body(doc, "审批通过率是消费金融经营的核心指标，但它同时受渠道结构、客群质量、审批规则和产品额度影响。只看整体通过率，容易把“结构变化”误判为“风控收紧”。")
    for item in [
        "问题一：通过率下降是渠道结构变化，还是同一渠道审批变严？",
        "问题二：哪些渠道贡献最大？",
        "问题三：通过率下降是否对应更好的风险表现？",
    ]:
        paragraph = doc.add_paragraph(style="List Bullet")
        run = paragraph.add_run(item)
        font(run, size=10)

    heading(doc, "三、数据与方法")
    body(doc, "数据来自 Home Credit 公开消费金融数据集，本次使用申请审批表与分期还款表。日期按 2026-09-30 为锚点映射，金额统一除以 10,000，单位为万元。")
    method = pd.DataFrame(
        [
            ("进件量", "全部申请笔数", "衡量规模"),
            ("审批通过率", "（Approved + Unused offer）/ 进件量", "衡量审批结果"),
            ("用信率", "Approved /（Approved + Unused offer）", "衡量批而不用"),
            ("结构效应", "渠道占比变化 × 基期渠道与整体通过率差", "衡量渠道组合变化"),
            ("比率效应", "当期渠道占比 × 渠道通过率变化", "衡量渠道自身变化"),
        ],
        columns=["指标/方法", "公式", "用途"],
    )
    table(doc, method, widths=[1.3, 3.7, 2.0])

    heading(doc, "四、趋势与漏斗")
    body(doc, f"{cur['月份']} 进件量 {fmt(cur['进件量'])} 笔，审批通过 {fmt(cur['审批通过量'])} 笔，实际用信 {fmt(cur['实际用信量'])} 笔，用信率 {pct(cur['用信率'])}。用信率处于高位，说明“批而不用”不是主要矛盾，矛盾集中在审批通过率。")
    doc.add_picture(str(figures / "01_月度趋势.png"), width=Inches(6.5))
    funnel = pd.DataFrame(
        [
            ("进件", fmt(cur["进件量"])),
            ("审批通过", fmt(cur["审批通过量"])),
            ("实际用信", fmt(cur["实际用信量"])),
            ("审批通过率", pct(cur["审批通过率"])),
            ("用信率", pct(cur["用信率"])),
        ],
        columns=["环节", "数值"],
    )
    table(doc, funnel, widths=[2.0, 2.0])

    heading(doc, "五、渠道归因")
    body(doc, f"整体通过率下降 {abs(total_effect)*100:.2f} 个百分点，其中结构效应 {mix_total*100:+.2f}pp、比率效应 {rate_total*100:+.2f}pp。结构效应是主因，但不能忽略比率效应。")
    top_negative = decomp.head(3).copy()
    top_negative["结构效应"] = top_negative["结构效应"].map(lambda x: f"{x*100:+.2f}pp")
    top_negative["比率效应"] = top_negative["比率效应"].map(lambda x: f"{x*100:+.2f}pp")
    top_negative["总效应"] = top_negative["总效应"].map(lambda x: f"{x*100:+.2f}pp")
    for _, row in top_negative.iterrows():
        body(
            doc,
            f"{row['渠道类型']}：结构 {row['结构效应']}、比率 {row['比率效应']}、合计 {row['总效应']}。",
            size=9.8,
        )
    doc.add_picture(str(figures / "03_渠道贡献.png"), width=Inches(6.4))

    heading(doc, "六、风险验证")
    body(doc, "把 2026-09 各渠道审批通过率与获批客户的 DPD30 放在一起看。注意：这里只能看到获批客户的风险，存在选择偏差，不能直接推断被拒绝客户的风险。")
    risk_table = risk_merged[["渠道类型", "审批通过率", "DPD30率"]].copy()
    risk_table["审批通过率"] = risk_table["审批通过率"].map(lambda x: pct(x))
    risk_table["DPD30率"] = risk_table["DPD30率"].map(lambda x: pct(x) if pd.notna(x) else "-")
    table(doc, risk_table, widths=[2.6, 1.4, 1.4])
    if (figures / "04_通过率与风险.png").exists():
        doc.add_picture(str(figures / "04_通过率与风险.png"), width=Inches(6.2))
    body(doc, "从图上看，审批通过率低的渠道并不必然对应更低 DPD30；部分低通过渠道的风险并不低。通过率下降不能简单解释为“风控更有效”，还需要结合渠道客群、拒绝原因和审批规则进一步验证。", size=9.8)

    heading(doc, "七、建议")
    for item in [
        "渠道管理部：对低通过率、高进件占比渠道做专项复核，拆解主动投放、被动流入和客群质量，验证指标为通过率、拒绝原因 Top3、单位获客成本。",
        "风险管理部：对通过率下降但 DPD30 未改善的渠道，检查审批规则与反欺诈策略，验证指标为通过率与 DPD30 联动变化。",
        "经营分析部：把结构/比率分解固化为月度监控，结构效应超过 2pp 时触发专项说明，避免把结构变化误读为风控能力变化。",
    ]:
        paragraph = doc.add_paragraph(style="List Bullet")
        run = paragraph.add_run(item)
        font(run, size=9.8)

    heading(doc, "八、局限")
    for item in [
        "公开数据、模拟业务日历，结论用于方法展示，不代表真实公司经营。",
        "金额为归一化万元，不能解释为真实人民币规模。",
        "风险验证只看获批客户，存在选择偏差，不能做因果推断。",
        "缺少渠道投放、拒绝原因、审批规则变更等业务事件数据。",
    ]:
        paragraph = doc.add_paragraph(style="List Bullet")
        run = paragraph.add_run(item)
        font(run, size=9.8)

    heading(doc, "九、附录：口径")
    body(doc, "审批通过率 =（Approved + Unused offer）/ 进件量；用信率 = Approved /（Approved + Unused offer）；DPD30 = 单期最大逾期天数 ≥ 30。所有数字均来自 MySQL 汇总表/视图，可回到 sql/ 目录复算。", size=9)

    docx_path = reports / f"审批通过率下降专项分析_{cur['月份']}.docx"
    doc.save(docx_path)
    pdf_path = docx_path.with_suffix(".pdf")
    script = (
        "$ErrorActionPreference='Stop'; "
        "$word=New-Object -ComObject Word.Application; $word.Visible=$false; "
        f"$doc=$word.Documents.Open('{docx_path}'); $doc.SaveAs([ref]'{pdf_path}',[ref]17); "
        "$doc.Close(); $word.Quit()"
    )
    try:
        subprocess.run(["powershell", "-NoProfile", "-Command", script], check=True, capture_output=True)
    except Exception as exc:
        print("PDF 转换失败（DOCX 已生成）:", exc)
    print(docx_path)
    print(pdf_path if pdf_path.exists() else "PDF 未生成")


if __name__ == "__main__":
    main()
